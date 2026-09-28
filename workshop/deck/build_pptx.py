from pptx import Presentation
from pptx.util import Inches
import os, re, io, slides as S

HERE = os.path.dirname(os.path.abspath(__file__))
GUIDE = os.path.abspath(os.path.join(HERE, "..", "facilitator", "Presenter-Guide_Part2_Slide-Notes.md"))

txt = io.open(GUIDE, encoding="utf8").read()
notes = {}
for m in re.finditer(r"## Slide (\d+) · ([^\n]+)\n(.*?)(?=\n## |\n# |\Z)", txt, re.S):
    notes[int(m.group(1))] = m.group(3).strip()

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]

for s in S.SLIDES:
    jpg = os.path.join(HERE, "jpg", S.filename_for(s).replace(".png", ".jpg"))
    sl = prs.slides.add_slide(blank)
    sl.shapes.add_picture(jpg, 0, 0, width=prs.slide_width, height=prs.slide_height)
    body = re.sub(r"\*\*(.*?)\*\*", r"\1", notes.get(s[0], ""))
    body = re.sub(r"^> ?", "", body, flags=re.M)
    head = "Slide %d: %s  [%s]" % (s[0], s[1].replace("-", " "), S.MODE_NAME[s[2]])
    sl.notes_slide.notes_text_frame.text = (head + "\n\n" + body)[:5000]

out = os.path.join(HERE, "Build-Your-90-Day-Plan_Deck.pptx")
prs.save(out)
print("pptx: %.1f MB, %d slides" % (os.path.getsize(out)/1e6, len(prs.slides._sldIdLst)))
